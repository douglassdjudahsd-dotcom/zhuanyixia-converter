from pathlib import Path
import struct, zlib, json, io
import pymupdf
from docx import Document

root=Path(__file__).resolve().parent.parent/'tests/fixtures'
root.mkdir(parents=True,exist_ok=True)
pdf=pymupdf.open()
for page_no in range(2):
    page=pdf.new_page(width=595,height=842)
    page.insert_text((60,80),'PDF conversion test %d'%(page_no+1),fontsize=18)
    page.insert_text((60,120),'Hello, editable Word. Amount: 12345.',fontsize=12)
    page.insert_text((60,160),'中文文件转换测试，保留可编辑文字。',fontname='china-s',fontsize=14)
    if page_no==1:
        for y in [220,255,290]: page.draw_line((60,y),(480,y))
        for x in [60,270,480]: page.draw_line((x,220),(x,290))
        page.insert_text((75,245),'Item A',fontsize=12)
        page.insert_text((285,245),'100',fontsize=12)
data=pdf.tobytes()
(root/'sample.pdf').write_bytes(data)
header=bytearray(512);header[:4]=b'CAJ\0';struct.pack_into('<i',header,0x10,2);struct.pack_into('<i',header,0x14,0x80);struct.pack_into('<i',header,0x80,512)
(root/'sample.caj').write_bytes(header+data)
parts=[]
for number in range(1,pdf.xref_length()):
    obj=pdf.xref_object(number,compressed=False)
    if '/Type /Catalog' in obj or '/Type /Pages' in obj: continue
    body=obj.encode('latin1')
    if pdf.xref_is_stream(number): body+=b'\nstream\n'+pdf.xref_stream_raw(number)+b'\nendstream'
    parts.append(str(number).encode()+b' 0 obj\n'+body+b'\nendobj\n')
(root/'missing-root.caj').write_bytes(header+b''.join(parts))
kdh=bytearray(254);kdh[:4]=b'KDH ';key=b'FZHMEI'
(root/'sample.kdh').write_bytes(kdh+bytes(b^key[i%6] for i,b in enumerate(data)))

text='CAJ文件转换测试\n正文是可编辑的中文。\n12345'
records=bytearray()
for char in text:
    if char=='\n': code=0xa38a
    else:
        enc=char.encode('gbk');code=int.from_bytes(enc,'big') if len(enc)==2 else enc[0]
    records.extend(struct.pack('<HHHH',0x8001,0,0,code))
compressed=b'COMPRESSTEXT'+struct.pack('<i',len(records))+zlib.compress(records)
hn=bytearray(512);hn[:6]=b'HN\0\0\xc8\0';struct.pack_into('<i',hn,0x90,1);struct.pack_into('<iihhii',hn,0xd8,512,len(compressed),0,1,0,0)
(root/'sample-hn.caj').write_bytes(hn+compressed)
c8=bytearray(512);c8[0]=0xc8;struct.pack_into('<i',c8,8,1);struct.pack_into('<iihhii',c8,0x50,512,len(records),0,1,0,0)
(root/'sample-c8.caj').write_bytes(c8+records)
old=bytearray()
for line in text.split('\n'):
    old.extend(struct.pack('<HH',0x8001,0))
    for char in line:
        enc=char.encode('gbk');code=int.from_bytes(enc,'big') if len(enc)==2 else enc[0]
        old.extend(struct.pack('<HH',0,code))
old_hn=bytearray(hn);struct.pack_into('<iihhii',old_hn,0xd8,512,len(old),0,1,0,512+len(old))
(root/'old-style-hn.caj').write_bytes(old_hn+old)
(root/'broken.caj').write_bytes(b'CAJ\0bad')

source=pymupdf.open(); p=source.new_page(width=750,height=475)
p.insert_text((50,82),'SCANNED TEST 12345',fontname='helv',fontsize=32)
p.insert_text((50,150),'文字识别测试',fontname='china-s',fontsize=32)
p.insert_text((50,217),'清晰的中英文扫描图片',fontname='china-s',fontsize=32)
p.get_pixmap(matrix=pymupdf.Matrix(2,2)).save(root/'scan.png')
scan=pymupdf.open();page=scan.new_page(width=750,height=475);page.insert_image(page.rect,stream=(root/'scan.png').read_bytes());scan.save(root/'scan.pdf',deflate=True)

document=Document();document.add_heading('Word转换测试',0);document.add_paragraph('Hello Word 12345，中文段落。');table=document.add_table(rows=2,cols=2);table.cell(0,0).text='姓名';table.cell(0,1).text='成绩';table.cell(1,0).text='小明';table.cell(1,1).text='90';document.save(root/'sample.docx')
(root/'sample.txt').write_text('文字转Word测试\nHello mobile.\n第二段中文 12345',encoding='utf-8')
(root/'sample.csv').write_text('姓名,成绩\n小明,90\n小红,95\n',encoding='utf-8-sig')
(root/'sample.json').write_text(json.dumps([{'姓名':'小明','成绩':90},{'姓名':'小红','成绩':95}],ensure_ascii=False),encoding='utf-8')
print('Created converter fixtures:',len(list(root.iterdir())))
